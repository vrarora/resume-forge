#!/usr/bin/env python3
"""Convert an ATS resume HTML (built from assets/templates/ats.html) to DOCX.

Usage:
  python3 to_docx.py resume.html "First Last Resume.docx" [--letter]

Needs python-docx (pip install python-docx). The DOCX uses Arial, because
Word users rarely have Geist installed. Page size defaults to A4.
"""

import sys
from html.parser import HTMLParser
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor
except ImportError:
    sys.exit("python-docx is missing. Run: pip install python-docx")

FONT = "Arial"
INK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x6B, 0x6B, 0x6B)
VOID = {"meta", "br", "img", "hr", "link", "input"}


# Minimal DOM

class Node:
    def __init__(self, tag, attrs=None, parent=None):
        self.tag, self.attrs, self.parent, self.children = tag, dict(attrs or {}), parent, []

    @property
    def classes(self):
        return self.attrs.get("class", "").split()

    def has(self, cls):
        return cls in self.classes

    def find_all(self, pred):
        out = []
        for c in self.children:
            if isinstance(c, Node):
                if pred(c):
                    out.append(c)
                out.extend(c.find_all(pred))
        return out


class Builder(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Node("root")
        self.cur = self.root
        self.skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("style", "script", "head", "title"):
            self.skip += 1
            return
        node = Node(tag, attrs, self.cur)
        self.cur.children.append(node)
        if tag not in VOID:
            self.cur = node

    def handle_endtag(self, tag):
        if tag in ("style", "script", "head", "title"):
            self.skip -= 1
            return
        n = self.cur
        while n is not self.root and n.tag != tag:
            n = n.parent
        if n is not self.root:
            self.cur = n.parent

    def handle_data(self, data):
        if not self.skip and data.strip():
            self.cur.children.append(data)


# Inline runs

def runs(node, bold=False):
    """Flatten a node into (text, href, bold) runs, collapsing whitespace."""
    out = []
    for c in node.children:
        if isinstance(c, str):
            out.append((" ".join(c.split()) if c.strip() else " ", None, bold))
        elif c.tag == "a":
            out.append((text_of(c), c.attrs.get("href"), bold))
        elif c.tag == "br":
            out.append(("\n", None, bold))
        else:
            out.extend(runs(c, bold or c.has("label") or c.has("title") and node.has("compact")))
    # Restore single spaces between adjacent runs
    fixed = []
    for i, (t, h, b) in enumerate(out):
        if i and fixed and not fixed[-1][0].endswith((" ", "\n")) and not t.startswith((" ", ",", ".", "\n")):
            t = " " + t
        fixed.append((t, h, b))
    return fixed


def text_of(node):
    return " ".join("".join(t for t, _, _ in runs(node)).split())


def add_hyperlink(par, text, url, size, color):
    rid = par.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for tag, val in (("w:rFonts", None), ("w:sz", str(int(size * 2))), ("w:u", "single"),
                     ("w:color", str(color))):
        el = OxmlElement(tag)
        if tag == "w:rFonts":
            el.set(qn("w:ascii"), FONT)
            el.set(qn("w:hAnsi"), FONT)
        else:
            el.set(qn("w:val"), val)
        rpr.append(el)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    link.append(r)
    par._p.append(link)


def write_runs(par, rs, size=9.5, color=INK):
    for text, href, bold in rs:
        if href:
            lead = text[: len(text) - len(text.lstrip())]
            if lead:
                write_runs(par, [(lead, None, bold)], size, color)
            add_hyperlink(par, text.strip(), href, size, color)
            continue
        r = par.add_run(text)
        r.bold = bold
        r.font.size = Pt(size)
        r.font.name = FONT
        r.font.color.rgb = color


# Document

class Doc:
    def __init__(self, letter=False):
        self.d = Document()
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = (Mm(215.9), Mm(279.4)) if letter else (Mm(210), Mm(297))
        sec.top_margin = sec.bottom_margin = Mm(14)
        sec.left_margin = sec.right_margin = Mm(17)
        self.width = sec.page_width - sec.left_margin - sec.right_margin
        st = self.d.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(9.5)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    def para(self, space_before=0, space_after=0):
        p = self.d.add_paragraph()
        pf = p.paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(space_before), Pt(space_after), 1.1
        return p

    def text(self, rs, size=9.5, color=INK, before=0, after=0):
        p = self.para(before, after)
        write_runs(p, rs, size, color)
        return p

    def row(self, left_rs, right, size=9.5, color=INK, before=0, bold_left=False):
        p = self.para(before, 0)
        p.paragraph_format.tab_stops.add_tab_stop(self.width, WD_TAB_ALIGNMENT.RIGHT)
        write_runs(p, [(t, h, b or bold_left) for t, h, b in left_rs], size, color)
        if right:
            r = p.add_run("\t" + right)
            r.font.size, r.font.name, r.font.color.rgb = Pt(8.5), FONT, GREY
        return p

    def bullet(self, rs):
        p = self.para(0, 1.5)
        pf = p.paragraph_format
        pf.left_indent, pf.first_line_indent = Mm(4), Mm(-4)
        pf.tab_stops.add_tab_stop(Mm(4))
        write_runs(p, [("•\t", None, False)] + rs)
        return p


def row_parts(node):
    spans = [c for c in node.children if isinstance(c, Node)]
    left = spans[0] if spans else node
    right = text_of(spans[-1]) if len(spans) > 1 else ""
    return left, right


def convert(src, dst, letter=False):
    b = Builder()
    b.feed(Path(src).read_text(encoding="utf-8"))
    body = (b.root.find_all(lambda n: n.tag == "body") or [b.root])[0]
    doc = Doc(letter)

    def walk(node):
        for c in node.children:
            if not isinstance(c, Node):
                continue
            if c.tag == "h1":
                doc.text([(text_of(c), None, False)], size=20)
            elif c.has("headline"):
                doc.text(runs(c), size=10.5, color=GREY, before=2)
            elif c.has("contact"):
                doc.text(runs(c), size=8.5, color=GREY, before=4)
            elif c.tag == "h2":
                doc.text([(text_of(c), None, True)], size=9.5, before=12, after=3)
            elif c.has("row"):
                left, right = row_parts(c)
                if left.has("co"):
                    doc.row([(text_of(left), None, True)], right, size=10.5, before=6)
                elif left.has("title"):
                    doc.row([(text_of(left), None, False)], right, color=GREY, before=3)
                else:
                    doc.row(runs(left), right)
            elif c.has("context") or c.has("awards") or c.has("note"):
                doc.text(runs(c), size=8.5, color=GREY, before=2 if c.has("awards") else 0)
            elif c.tag == "ul":
                for li in c.children:
                    if isinstance(li, Node) and li.tag == "li":
                        doc.bullet(runs(li))
            elif c.tag == "p":
                doc.text(runs(c), size=9 if node.has("small") else 9.5, after=1)
            else:
                walk(c)

    walk(body)
    doc.d.save(dst)
    print(f"{Path(dst).name}: written")


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not args:
        sys.exit(__doc__)
    out = args[1] if len(args) > 1 else str(Path(args[0]).with_suffix(".docx"))
    convert(args[0], out, letter="--letter" in sys.argv)
